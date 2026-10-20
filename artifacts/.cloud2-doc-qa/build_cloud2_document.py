from pathlib import Path
from collections import Counter, defaultdict
import re
import json
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).parent
source = ROOT / '76.129-cloud2-image-inventory.md'
text = source.read_text(encoding='utf-8')
left, right = text.split('## 容器对应关系')
images = []
for line in left.splitlines():
    if line.startswith('| harbor.'):
        ref, iid, size, usage = [v.strip() for v in line.split('|')[1:-1]]
        repo, tag = ref.rsplit(':', 1)
        images.append(dict(ref=ref, repo=repo, tag=tag, name=repo.rsplit('/', 1)[-1], iid=iid, size=size, usage=usage))
containers = []
for line in right.splitlines():
    if line.startswith('| ') and 'harbor.' in line:
        name, ref, status = [v.strip() for v in line.split('|')[1:-1]]
        containers.append(dict(name=name, ref=ref, status=status, running=status.startswith('Up ')))
assert len(images) == 128 and len(containers) == 98
image_by_ref = {i['ref']: i for i in images}
running = [c for c in containers if c['running']]
exited = [c for c in containers if not c['running']]
active = defaultdict(list)
for c in running:
    active[c['ref']].append(c)
assert len(active) == 82

# Descriptions are an operational reading of component names, not verified API contracts.
roles = {
'uranus-fe': ('业务与设备服务','平台前端页面'),
'pf-gateway-service': ('业务与设备服务','平台网关与对外接口入口'),
'pf-device-service': ('业务与设备服务','设备管理'),
'pf-device-simulator-service': ('业务与设备服务','设备模拟'),
'pf-control-core': ('业务与设备服务','设备控制核心'),
'es-drc-service': ('业务与设备服务','DRC 远程控制相关服务'),
'es310-management-service': ('业务与设备服务','平台业务管理'),
'es-task-management': ('业务与设备服务','任务管理'),
'file-storage': ('业务与设备服务','文件存储业务接口'),
'pf-notify-service': ('业务与设备服务','消息通知'),
'pf-inventory-collector': ('业务与设备服务','资源清单采集'),
'pf-terra-manage-service': ('业务与设备服务','Terra 建模任务管理'),
'geo-resource-service': ('业务与设备服务','地理资源管理'),
'pf-devkit-mcp': ('业务与设备服务','开发套件 MCP 接口'),
'thing-access-service': ('业务与设备服务','物联网设备接入'),
'es-flysafe': ('业务与设备服务','飞行安全相关服务'),
'pf-live-service': ('直播与实时通信','平台侧直播业务管理'),
'live-service': ('直播与实时通信','直播服务'),
'live-signaling': ('直播与实时通信','直播信令'),
'stream-hub': ('直播与实时通信','直播流汇聚与分发'),
'srs-custom-srs': ('直播与实时通信','SRS 流媒体处理'),
'zlmediakit': ('直播与实时通信','ZLMediaKit 流媒体处理'),
'live-snapshot': ('直播与实时通信','直播截图'),
'drtc-manager': ('直播与实时通信','实时通信管理'),
'lambda-audio-convert': ('航线媒体与算法处理','音频转换'),
'lambda-ppe': ('航线媒体与算法处理','PPE 算法执行组件'),
'serverless-wayline-process': ('航线媒体与算法处理','航线文件处理'),
'serverless-wayline-zip': ('航线媒体与算法处理','航线打包'),
'serverless-wayline-unzip': ('航线媒体与算法处理','航线解包'),
'serverless-flight-replay-parse': ('航线媒体与算法处理','飞行回放数据解析'),
'serverless-image-callback': ('航线媒体与算法处理','图片处理回调'),
'serverless-image-draw': ('航线媒体与算法处理','图片绘制'),
'serverless-video-callback': ('航线媒体与算法处理','视频处理回调'),
'serverless-video-snapshot': ('航线媒体与算法处理','视频截图'),
'serverless-ai-video': ('航线媒体与算法处理','AI 视频处理'),
'serverless-html2file': ('航线媒体与算法处理','HTML 转文件'),
'lambda': ('航线媒体与算法处理','通用函数运行时；职责由容器加载的函数确定'),
'algorithm-ktx2': ('航线媒体与算法处理','KTX2 纹理处理'),
'algorithm-model': ('航线媒体与算法处理','模型处理'),
'las-to-pnts': ('航线媒体与算法处理','LAS 点云转 PNTS'),
'tif-to-terrain': ('航线媒体与算法处理','TIF 转地形数据'),
'terra-rescon-be': ('航线媒体与算法处理','Terra 重建相关后端'),
'offline-map-lambda': ('地图与高程服务','离线地图处理'),
'digital-elevation-data-service': ('地图与高程服务','数字高程数据服务'),
'mbtileserver': ('地图与高程服务','MBTiles 地图瓦片服务'),
'mysql': ('数据存储消息与入口','MySQL 关系数据库'),
'postgres': ('数据存储消息与入口','PostgreSQL 关系数据库'),
'postgis': ('数据存储消息与入口','PostgreSQL 与 PostGIS 空间数据库镜像'),
'timescaledb': ('数据存储消息与入口','TimescaleDB 时序数据库镜像'),
'mongo': ('数据存储消息与入口','MongoDB 文档数据库'),
'redis': ('数据存储消息与入口','Redis 内存数据存储'),
'rabbitmq': ('数据存储消息与入口','RabbitMQ 消息队列'),
'emqx': ('数据存储消息与入口','EMQX MQTT 消息代理'),
'emqx-broker': ('数据存储消息与入口','设备接入相关 MQTT 消息代理'),
'minio': ('数据存储消息与入口','MinIO 对象存储'),
'nginx': ('数据存储消息与入口','HTTP 服务与反向代理；对应多个容器角色'),
'traefik': ('数据存储消息与入口','HTTP 路由与反向代理'),
'ntp': ('数据存储消息与入口','NTP 时间同步'),
'prometheus': ('监控调度与运维','指标采集与存储'),
'grafana': ('监控调度与运维','指标与日志可视化'),
'alertmanager': ('监控调度与运维','告警汇总与路由'),
'loki': ('监控调度与运维','日志聚合存储'),
'promtail': ('监控调度与运维','日志采集'),
'pushgateway': ('监控调度与运维','接收推送指标'),
'node-exporter': ('监控调度与运维','主机指标导出'),
'cadvisor': ('监控调度与运维','容器资源指标采集'),
'docker-stats-exporter': ('监控调度与运维','Docker 容器统计指标导出'),
'blackbox-exporter': ('监控调度与运维','外部探测指标导出'),
'mysqld-exporter': ('监控调度与运维','MySQL 指标导出'),
'postgres-exporter': ('监控调度与运维','PostgreSQL 指标导出'),
'mongodb_exporter': ('监控调度与运维','MongoDB 指标导出'),
'redis_exporter': ('监控调度与运维','Redis 指标导出'),
'emqx-exporter': ('监控调度与运维','EMQX 指标导出'),
'dronenet-monitor': ('监控调度与运维','设备网络监测'),
'portainer-ce': ('监控调度与运维','容器运维界面'),
'es-privatization-system-be': ('监控调度与运维','私有化系统管理后端'),
'es310-privatization-devops-be': ('监控调度与运维','私有化运维后端'),
'es310-privatization-devops-fe': ('监控调度与运维','私有化运维前端'),
'es-task-admin': ('监控调度与运维','任务调度管理'),
'es-task-gateway': ('监控调度与运维','任务调度网关'),
'es-task-scheduler': ('监控调度与运维','任务调度执行管理'),
}
groups = ['业务与设备服务','直播与实时通信','航线媒体与算法处理','地图与高程服务','数据存储消息与入口','监控调度与运维']
intro = {
groups[0]: '这组服务覆盖前端、网关、设备、控制、任务、文件、通知和地理资源，是理解基座业务划分的主要入口。设备接入镜像来自 djiservice 命名空间，其余主要来自 es310 命名空间。',
groups[1]: '直播业务管理、信令和媒体流处理采用独立镜像。pf-live-service、live-service、SRS、ZLMediaKit 与 stream-hub 同时运行；组件之间的实际流向需要结合运行配置确认。',
groups[2]: '航线、音视频和三维数据处理主要采用 lambda 或 serverless 组件。lambda-ppe 有 5 个运行容器，lambda-audio-convert 有 3 个；容器数量代表当前部署实例数，不能直接等同于并发能力。',
groups[3]: '地图与高程相关组件包含瓦片服务、离线地图处理和高程数据服务。map-server-nginx 复用了公共 Nginx 镜像，在基础组件表中列出。',
groups[4]: '本机同时运行多类数据库、消息代理和对象存储。镜像版本可直接用于版本核对；具体数据库业务归属、存储卷和访问方式应以部署配置为准。',
groups[5]: '监控包括指标、日志、探测和多种 exporter。调度组件与私有化运维服务单独部署，分别承担任务调度和平台运行管理相关职责。',
}
group_rows = defaultdict(list)
for ref, cs in sorted(active.items(), key=lambda kv: kv[1][0]['name']):
    im = image_by_ref[ref]
    assert im['name'] in roles, im['name']
    group, role = roles[im['name']]
    group_rows[group].append([ '\n'.join(c['name'] for c in cs), im['name'] + '\n' + im['tag'], role, str(len(cs)) ])

doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Cm(21), Cm(29.7)
section.top_margin, section.bottom_margin = Cm(1.7), Cm(1.6)
section.left_margin, section.right_margin = Cm(1.6), Cm(1.6)
section.footer_distance = Cm(.7)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3']:
    s = doc.styles[name]
    s.font.name = 'Arial'
    s.font.color.rgb = RGBColor(0, 0, 0)
    s.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
doc.styles['Normal'].font.size = Pt(10)
doc.styles['Normal'].paragraph_format.space_after = Pt(6)
doc.styles['Normal'].paragraph_format.line_spacing = 1.15
doc.styles['Title'].font.size = Pt(23)
doc.styles['Title'].font.bold = True
doc.styles['Heading 1'].font.size = Pt(15)
doc.styles['Heading 1'].paragraph_format.space_before = Pt(14)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(7)
doc.styles['Heading 2'].font.size = Pt(12)
doc.styles['Heading 2'].paragraph_format.space_before = Pt(12)
doc.styles['Heading 2'].paragraph_format.space_after = Pt(6)
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = footer.add_run()
r.font.size = Pt(8)
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); r._r.addnext(field)
doc.core_properties.title = '上云2基座镜像与组件说明'
doc.core_properties.subject = '192.168.76.129 容器与镜像资产说明'
doc.core_properties.author = ''
md = []

def title(t):
    p = doc.add_paragraph(t, 'Title')
    borders = OxmlElement('w:pBdr')
    for edge in ['top','left','bottom','right','between','bar']:
        border = OxmlElement('w:'+edge); border.set(qn('w:val'),'nil'); borders.append(border)
    p._p.get_or_add_pPr().append(borders)
    md.extend(['# ' + t, ''])

def heading(t, level=1, page=False):
    if page: doc.add_page_break()
    doc.add_heading(t, level=level); md.extend(['#' * (level+1) + ' ' + t, ''])

def para(t):
    doc.add_paragraph(t); md.extend([t, ''])

def table(headers, rows, widths, small=False):
    tab = doc.add_table(rows=1, cols=len(headers))
    tab.alignment = WD_TABLE_ALIGNMENT.CENTER
    tab.autofit = False
    for c, w in zip(tab.columns, widths): c.width = Cm(w)
    for idx, row_vals in enumerate([headers] + rows):
        row = tab.rows[0] if idx == 0 else tab.add_row()
        trpr = row._tr.get_or_add_trPr()
        cant = OxmlElement('w:cantSplit'); trpr.append(cant)
        if idx == 0:
            repeat = OxmlElement('w:tblHeader'); trpr.append(repeat)
        for j, (cell, val, w) in enumerate(zip(row.cells, row_vals, widths)):
            cell.width = Cm(w)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            borders = OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                b = OxmlElement('w:'+edge); b.set(qn('w:val'),'single'); b.set(qn('w:sz'),'4'); b.set(qn('w:color'),'D9D9D9'); borders.append(b)
            tcpr.append(borders)
            margins = OxmlElement('w:tcMar')
            for edge, valm in [('top','65'),('bottom','65'),('left','85'),('right','85')]:
                m = OxmlElement('w:'+edge); m.set(qn('w:w'),valm); m.set(qn('w:type'),'dxa'); margins.append(m)
            tcpr.append(margins)
            shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'DCE6F1' if idx == 0 else ('F5F7FA' if idx%2 == 0 else 'FFFFFF')); tcpr.append(shade)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.06
            p.paragraph_format.keep_with_next = idx == 0
            if j == len(headers)-1 and headers[-1] in ['实例','状态','数量']: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(str(val))
            run.font.name = 'Arial'; run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
            run.font.size = Pt(8 if small else 8.5); run.bold = idx == 0
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    md.append('| ' + ' | '.join(headers) + ' |')
    md.append('| ' + ' | '.join('---' for _ in headers) + ' |')
    for row in rows: md.append('| ' + ' | '.join(str(v).replace('\n','<br>') for v in row) + ' |')
    md.append('')

title('上云2基座镜像与组件说明')
para('服务器 192.168.76.129')
para('本说明面向基座对接开发和日常运维，介绍本机已发现的业务服务、算法处理与基础组件，并汇总镜像版本和容器状态。当前快照包含 128 个镜像标签、98 个容器，其中 93 个容器运行中，使用 82 个不同镜像引用。')
heading('一 环境与统计口径')
table(['项目','记录'],[
['目标主机','192.168.76.129'],['部署目录','/opt/cloudApi2_3.1'],['主要编排文件','/opt/cloudApi2_3.1/install/compose.yaml'],['镜像仓库','harbor.djicorp.com'],['数据来源','本次会话已采集的 docker image ls 与 docker ps -a 快照'],['文档日期','2026年10月19日'],['用途','组件识别、镜像版本核对与维护交接']],[4,13.8])
table(['统计项','数量'],[['本地镜像标签',128],['全部容器',98],['运行中容器',93],['已退出容器',5],['运行中容器使用的不同镜像引用',82],['仅已退出容器引用的镜像标签',2],['无现有容器引用的镜像标签',44]],[13.8,4])
para('这里的“镜像”按仓库与标签组成的完整引用统计；容器可以复用同一个镜像。128 = 82 + 2 + 44，98 = 93 + 5。镜像显示大小不能直接相加作为磁盘占用，因为镜像之间可能共享数据层。')
para('容器状态中，43 个显示 healthy，50 个显示 Up 且未附带健康状态，5 个显示 Exited (0)。Up 只表示容器进程运行，未附带 healthy 不能据此判定异常；healthy 也不等同于完整业务链路验收通过。')
heading('二 组件分类与运行版本')
para('以下按运行中容器使用的 82 个不同镜像引用分组，覆盖全部 93 个运行实例。镜像列第一行是镜像名，第二行是标签；完整仓库路径、镜像 ID 和大小见附录。通用开源组件采用其常见职责说明，私有服务职责按组件名称作概括性解读，精确接口和调用关系以部署配置及产品接口文档为准。')
table(['组件分类','不同镜像','运行实例'],[[g,str(len(group_rows[g])),str(sum(int(r[-1]) for r in group_rows[g]))] for g in groups],[10.2,3.8,3.8])
for idx,g in enumerate(groups,1):
    heading(f'{idx} {g}',2)
    para(intro[g])
    table(['容器名称','镜像名与标签','职责概览','实例'],group_rows[g],[5.1,6.1,5.6,1])
    if g == '航线媒体与算法处理':
        para('公共 lambda:go1.x 镜像由图片裁剪、视频转换和 PSDK 喊话器音频转码三个容器复用；lambda:python3.8 对应 lambda-zip-file。镜像相同不代表加载的函数相同。PPE 的缩写及算法边界需以厂商说明为准。')
    if g == '数据存储消息与入口':
        para('minio 与 minio-static 共享 MinIO 镜像；forbid-service、lambda-nginx 和 map-server-nginx 共享 Nginx 镜像。EMQX 的 4.3.5 与 emqx-broker 的 3323811 来自不同镜像仓库，分别保留原始标签，不能按同一版本序列比较。')

heading('三 已退出容器与未引用镜像')
para('5 个已退出容器均显示退出码 0。名称表明它们涉及数据库初始化、对象存储建桶或高程更新；这种状态与一次性任务完成相符，不能仅凭退出状态判定部署故障。')
table(['容器','镜像名与标签','快照状态'],[[c['name'],image_by_ref[c['ref']]['name']+'\n'+image_by_ref[c['ref']]['tag'],c['status']] for c in exited],[6.7,6.6,4.5])
para('mc:latest 和 elevation-data-update-tool:2805567 这两个镜像标签仅被已退出容器引用。MySQL 与 PostGIS 镜像同时用于数据库运行容器和初始化容器，因此仍计入运行中使用的 82 个镜像引用。')
unused = [i for i in images if i['usage'].startswith('无现有')]
active_repos = {image_by_ref[ref]['repo'] for ref in active}
alternate = [i for i in unused if i['repo'] in active_repos]
other = [i for i in unused if i['repo'] not in active_repos]
para(f'另有 {len(unused)} 个镜像标签未被快照中的现有容器引用。其中 {len(alternate)} 个与运行中的镜像具有相同仓库但标签不同，可作为已有服务的其他版本记录；其余 {len(other)} 个属于快照中没有运行容器的镜像仓库。没有现有容器引用不等同于从未使用，也不等同于可以删除。')
heading('现有服务的其他本地版本',2)
alts = defaultdict(list)
for i in alternate: alts[i['repo']].append(i['tag'])
table(['镜像','运行中标签','其他本地标签'],[[repo.rsplit('/',1)[-1], '\n'.join(image_by_ref[r]['tag'] for r in active if image_by_ref[r]['repo']==repo),'\n'.join(sorted(tags))] for repo,tags in sorted(alts.items())],[7,4.3,6.5])
heading('未发现运行容器的镜像用途',2)
para('这一部分包含 pf-atc-service、pf-wayline-service、pf-media-core、pf-safety-service，以及 AEC、报告、全景、模型截图、KML 导入等算法或函数镜像。还包括 kubectl 和 NVIDIA 设备插件镜像。它们的本地存在只能证明镜像已缓存；按需任务、历史容器、其他编排运行时或 GPU 实际启用情况不能由本快照确认。全部标签及大小见附录。')

heading('四 版本核对与维护参考')
para('核对某个组件时，先按容器名称找到正在使用的完整镜像引用，再对照附录中的 ID 与大小。数字构建号、带 _s 的标签、语义版本号和时间戳标签均按原样保留；不得把不同组件的标签数值大小解释为统一的发布先后关系。目录名 cloudApi2_3.1 表示本机部署目录命名，不能单独作为产品发布版本的证明。')
para('MinIO、mc、NTP 和 kubectl 使用 latest 标签。后续做可复现交付时，建议同时记录 RepoDigest；本清单中的短镜像 ID 可用于本机识别，但不替代跨环境拉取时使用的摘要。')
para('如需清理未引用镜像，应先核对安装包、升级回滚要求、按需函数和其他运行时的引用情况。本文件仅记录资产与状态，不包含镜像删除、组件升级或服务重启操作。')
heading('只读复核命令',2)
commands = [
'sudo docker image ls --digests --no-trunc',
'sudo docker ps -a --format "{{.Names}} | {{.Image}} | {{.Status}}"',
'sudo docker image inspect <image:tag> --format "{{json .RepoDigests}}"',
'sudo docker system df -v',
]
for cmd in commands:
    p = doc.add_paragraph(); p.paragraph_format.space_after=Pt(4)
    rr=p.add_run(cmd);rr.font.name='Consolas';rr.font.size=Pt(8.5)
md.extend(['```bash',*commands,'```',''])
para('复核命令适用于有相应读取权限的服务器会话。统计口径仍以同一次采集的镜像列表和容器列表为准，避免把不同时间点的数据直接比较。')

heading('附录 完整镜像标签清单')
para('本附录逐项保留 128 条镜像记录。为便于排版，仓库前缀用 P1 至 P11 表示，完整引用的还原方式为“前缀 + 镜像名 + 冒号 + 标签”。同一单元格中镜像名位于第一行，标签位于后续行。')
prefixes = sorted({i['repo'].rsplit('/',1)[0]+'/' for i in images})
prefix_code = {p:f'P{k}' for k,p in enumerate(prefixes,1)}
table(['代码','镜像仓库前缀'],[[prefix_code[p],p] for p in prefixes],[1.8,16])
para('状态“运行”表示至少一个运行中容器引用该镜像；“退出”表示只有已退出容器引用；“未引用”表示快照中的现有容器均未引用该完整标签。短镜像 ID 与大小均保留原始采集值。')
inv_rows=[]
for i in images:
    prefix=i['repo'].rsplit('/',1)[0]+'/'
    state='运行' if i['usage']=='运行中' else '退出' if i['usage']=='仅已退出容器引用' else '未引用'
    inv_rows.append([prefix_code[prefix],i['name']+'\n'+i['tag'],i['iid'],i['size'],state])
table(['前缀','镜像名与标签','短镜像 ID','大小','状态'],inv_rows,[1.2,10,2.7,1.7,2.2],small=True)

outdoc=ROOT/'上云2基座镜像与组件说明.docx'
outmd=ROOT/'上云2基座镜像与组件说明.md'
doc.save(outdoc)
# Markdown retains full references in the inventory, making exact text lookup simple.
start=md.index('## 附录 完整镜像标签清单')
md=md[:start]+['## 附录 完整镜像标签清单','','| 完整镜像引用 | 短镜像 ID | 大小 | 容器引用状态 |','| --- | --- | --- | --- |']
for i in images: md.append(f"| `{i['ref']}` | `{i['iid']}` | {i['size']} | {i['usage']} |")
outmd.write_text('\n'.join(md)+'\n',encoding='utf-8')
assert len(inv_rows)==128
assert sum(sum(int(r[-1]) for r in group_rows[g]) for g in groups)==93
assert len({c['name'] for c in containers})==98
print(json.dumps({'docx':str(outdoc),'markdown':str(outmd),'images':len(images),'running_containers':len(running),'running_image_refs':len(active),'alternate_version_tags':len(alternate),'other_unreferenced_tags':len(other),'prefixes':prefix_code},ensure_ascii=False,indent=2))
