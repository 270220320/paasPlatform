import sys,json
from pathlib import Path
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT

creds=json.load(sys.stdin)
doc=Document();s=doc.sections[0]
s.page_width=Cm(21);s.page_height=Cm(29.7)
s.top_margin=Cm(1.8);s.bottom_margin=Cm(1.8);s.left_margin=Cm(1.7);s.right_margin=Cm(1.7)
for name in ['Normal','Title','Heading 1']:
 st=doc.styles[name];st.font.name='Arial';st.font.color.rgb=RGBColor(0,0,0)
 st.element.rPr.rFonts.set(qn('w:eastAsia'),'微软雅黑')
 st.paragraph_format.space_after=Pt(8)
for b in list(doc.styles.element.iter(qn('w:pBdr'))):b.getparent().remove(b)
doc.styles['Normal'].font.size=Pt(10.5);doc.styles['Normal'].paragraph_format.line_spacing=1.15
doc.styles['Title'].font.size=Pt(23);doc.styles['Heading 1'].font.size=Pt(14)
doc.add_paragraph('上云2基座数据库账号与连接信息','Title')
doc.add_paragraph('目标服务器 192.168.76.129')
doc.add_paragraph('本文记录已在本机 Navicat 中验证成功的四套数据库登录凭据。四个账号均具备写入和管理权限；同一实例下的业务库使用对应实例的账号登录。')

def table(headers,rows,widths,password_column=None):
 t=doc.add_table(rows=1,cols=len(headers));t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
 for col,w in zip(t.columns,widths):col.width=Cm(w)
 for cell,value in zip(t.rows[0].cells,headers):cell.text=value
 t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
 for row in rows:
  for cell,value in zip(t.add_row().cells,row):cell.text=value
 for i,row in enumerate(t.rows):
  row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
  for j,cell in enumerate(row.cells):
   cell.width=Cm(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   pr=cell._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');b.append(el)
   pr.append(b);m=OxmlElement('w:tcMar')
   for edge in ['top','bottom','left','right']:
    el=OxmlElement('w:'+edge);el.set(qn('w:w'),'110');el.set(qn('w:type'),'dxa');m.append(el)
   pr.append(m);fill=OxmlElement('w:shd');fill.set(qn('w:fill'),'DDE8F1' if i==0 else ('F5F7F9' if i%2==0 else 'FFFFFF'));pr.append(fill)
   for par in cell.paragraphs:
    par.paragraph_format.space_after=Pt(1);par.paragraph_format.line_spacing=1.1
    for r in par.runs:
     r.font.size=Pt(9.5);r.bold=i==0
     if i and j==password_column:r.font.name='Consolas'
 doc.add_paragraph()

items=[('MySQL 8','mysql_instance_1','3306','root',''),('PostgreSQL 16','postgres_instance_1','5432','app_rw','postgres'),('PostGIS','postgres_instance_2','5432','app_rw','postgres'),('TimescaleDB','timescaledb','5432','postgres','postgres')]
doc.add_paragraph('数据库登录凭据','Heading 1')
rows=[]
for label,key,port,user,db in items:
 c=creds[key];assert c['user']==user and c['password']
 rows.append([label,c['ip']+':'+port,c['user'],c['password']])
table(['数据库实例','数据库内网地址','用户名','密码'],rows,[3.1,3.3,2.1,9.1],password_column=3)
doc.add_paragraph('密码区分大小写，复制时应完整保留字符。上述地址是容器内网地址，Navicat 使用 SSH 隧道访问。')
doc.add_paragraph('SSH 隧道与初始数据库','Heading 1')
doc.add_paragraph('SSH 主机：192.168.76.129　端口：22　用户名：test。SSH 登录使用本机 Navicat 已保存的配置。')
table(['数据库实例','Navicat 初始数据库','账号权限'],[
 ['MySQL 8','留空，连接后选择业务库','全局管理权限'],
 ['PostgreSQL 16','postgres','超级用户'],
 ['PostGIS','postgres','超级用户'],
 ['TimescaleDB','postgres','超级用户'],
],[4.0,6.3,7.3])
doc.add_paragraph('本机连接名称均以“76.129 上云2”开头。容器重建后应重新确认内网 IP；数据库密码发生变更后，应同步更新 Navicat 和本文档。')
doc.core_properties.title='上云2基座数据库账号与连接信息';doc.core_properties.subject='192.168.76.129 四套数据库登录凭据';doc.core_properties.author=''
out=Path('D:/AiFlow/artifacts/上云2基座数据库账号与连接信息.docx');doc.save(out)
check=Document(out)
for row,(_,key,_,_,_) in zip(check.tables[0].rows[1:],items):
 assert row.cells[3].text==creds[key]['password']
print('Document created and all four credentials verified')
