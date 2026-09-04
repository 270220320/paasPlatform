# -*- coding: utf-8 -*-
"""
更新 KOC3 报告需求第 6 条详细内容（配合用户已改的标题"限飞区与电子围栏在航线编辑中的展示"），
并重排需求表序号为连续的 1..N。
"""
import sys
from docx import Document

sys.stdout.reconfigure(encoding='utf-8')

PATH = r"D:\KOC3\大疆上云 API 2 KOC 3 测试报告模板_20260826(1).docx"

NEW_CONTENT = (
    "业务痛点：航线规划阶段无法在航线编辑器内查看限飞区/禁飞区与作业电子围栏，规划存在盲区，"
    "航线是否侵入受限空域只能在规划完成后人工比对或后端事后校验，安全风险无法前置；"
    "目前需在自有平台地图上单独叠加限飞区图层辅助判断。"
    "需求内容：航线编辑器（WaylineEditor）支持叠加展示官方限飞区/禁飞区图层与自定义电子围栏（作业区边框），"
    "航点或航线段进入受限区域时实时高亮提示，并支持将航线约束在围栏范围内规划，从源头避免违规航线。"
)

def set_cell_text(cell, text):
    for p in cell.paragraphs[1:]:
        p._element.getparent().remove(p._element)
    p = cell.paragraphs[0]
    runs = p.runs
    if not runs:
        p.add_run(text)
    else:
        runs[0].text = text
        for r in runs[1:]:
            r._element.getparent().remove(r._element)

def main():
    doc = Document(PATH)
    t2 = doc.tables[2]  # 需求表

    # 1) 更新第 6 条详细内容
    updated = False
    for r in t2.rows[1:]:
        if r.cells[1].text.strip() == '限飞区与电子围栏在航线编辑中的展示':
            set_cell_text(r.cells[2], NEW_CONTENT)
            updated = True
    assert updated, "未找到目标行"

    # 2) 序号重排为 1..N
    for i, r in enumerate(t2.rows[1:], start=1):
        set_cell_text(r.cells[0], str(i))

    doc.save(PATH)
    print("更新完成：第 6 条内容已改写，序号已重排为 1..%d" % (len(t2.rows) - 1))

if __name__ == "__main__":
    main()
