# -*- coding: utf-8 -*-
"""生成《关于推行"AI人机协同"开发流程的建议书》Word 文档"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn


def set_font(run, name_cn="宋体", name_en="Times New Roman", size=12,
             bold=False, color=None):
    run.font.name = name_en
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name_cn)
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = color


def add_title(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    set_font(r, "黑体", "SimHei", 18, bold=True)
    p.paragraph_format.space_after = Pt(12)


def add_subtitle(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    set_font(r, "楷体", "KaiTi", 14, bold=False)
    p.paragraph_format.space_after = Pt(18)


def add_h1(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_font(r, "黑体", "SimHei", 15, bold=True)
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(8)


def add_h2(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_font(r, "黑体", "SimHei", 13, bold=True)
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(6)


def add_para(doc, text, indent=True, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    p.paragraph_format.line_spacing = 1.4
    if indent:
        p.paragraph_format.first_line_indent = Pt(24)
    if bold_prefix:
        r1 = p.add_run(bold_prefix)
        set_font(r1, "宋体", "Times New Roman", 12, bold=True)
    r = p.add_run(text)
    set_font(r, "宋体", "Times New Roman", 12)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    p.paragraph_format.line_spacing = 1.4
    p.paragraph_format.left_indent = Pt(24)
    r = p.add_run(text)
    set_font(r, "宋体", "Times New Roman", 12)
    return p


def add_table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, h in enumerate(headers):
        cell = t.rows[0].cells[j]
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = cell.paragraphs[0].add_run(h)
        set_font(r, "黑体", "SimHei", 11, bold=True)
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            cell = t.rows[i + 1].cells[j]
            cell.paragraphs[0].alignment = (
                WD_ALIGN_PARAGRAPH.CENTER if j > 0 else WD_ALIGN_PARAGRAPH.LEFT)
            r = cell.paragraphs[0].add_run(v)
            set_font(r, "宋体", "Times New Roman", 11)
    if widths:
        for j, w in enumerate(widths):
            for row in t.rows:
                row.cells[j].width = Cm(w)
    doc.add_paragraph()
    return t


doc = Document()
# 页边距
for sec in doc.sections:
    sec.top_margin, sec.bottom_margin = Cm(2.54), Cm(2.54)
    sec.left_margin, sec.right_margin = Cm(3.0), Cm(3.0)

add_title(doc, "关于推行“AI人机协同”开发流程的建议")
add_subtitle(doc, "——基于无人机上云平台项目实践，对比KOC2传统编码模式的思考")

# ==================== 一、现状描述 ====================
add_h1(doc, "一、现状描述")

add_h2(doc, "（一）部门开发模式现状")
add_para(doc, "部门长期承担无人机上云类软件项目的研发与交付工作，需要深度对接大疆司空2（KOC2）"
              "及其FlightHub生态。在此类项目的早期建设中，我们采用传统编码模式：需求评审、人工设计、"
              "逐行手写代码、人工编写测试与排查脚本、上线后凭经验定位问题。这一模式在项目规模小、"
              "对接协议简单时尚可支撑，但随着SDK能力面扩大、交付节奏加快，其局限性日益凸显。")

add_h2(doc, "（二）传统编码模式存在的主要问题")
add_para(doc, "1．SDK知识消化成本高，调研周期长。", bold_prefix="痛点一：")
add_para(doc, "以司空2对接为例，FlightHub SDK覆盖航线、任务、设备、媒体、直播、HMS告警等十余个模块，"
              "错误码数千条，且存在大量“文档未写明、只能靠试错验证”的隐性行为（如媒体folderUuid的获取链路、"
              "配置在YML与数据库之间的优先级关系等）。靠人工通读文档加试错，单个功能的调研普遍需要3～5天，"
              "复杂功能（如媒体上传自治）超过一周。")

add_para(doc, "2．线上排障依赖“脚本海”战术，同一问题反复重试。", bold_prefix="痛点二：")
add_para(doc, "对项目库的统计可以直观反映这一成本：仅司空2对接期间积累的自动化诊断脚本就达87个，"
              "后端日志文件108个。更值得注意的是同一问题需要多轮重写脚本——机场设备绑定问题连续编写了"
              "10版诊断脚本（diag-dock3-bind1～10）、云服务切换问题9版、媒体上传链路问题15版以上、"
              "单个飞行任务的排障（task-314004）也迭代了15版脚本。一个线上问题平均需要人工编写5～10个"
              "排查脚本、跨越1～3天才能闭环，排查高度依赖个别“熟手”，人力被大量消耗在重复劳动上。")

add_para(doc, "3．知识沉淀在个人头脑中，组织能力存在断档风险。", bold_prefix="痛点三：")
add_para(doc, "大量踩坑经验（如FlexNet网关地址改写、KOC3鉴权方式由app_key调整为user_id、"
              "固件夜间挂死与网络配置的关联等）以口口相传和个人笔记的形式存在，没有转化为组织级可复用资产。"
              "一旦骨干人员流动，同类问题需要从头再踩一遍坑，返工成本高。")

add_para(doc, "4．样板代码占比高，挤占核心设计时间。", bold_prefix="痛点四：")
add_para(doc, "CRUD接口、DTO映射、配置类、前端表单页面等重复性代码约占日常编码工作量的60%，"
              "开发人员大量时间耗在“搬运”而非业务设计与质量打磨上，代码评审与自测时间被压缩，"
              "进一步推高了缺陷率与返工率。")

add_h2(doc, "（三）已有的转型实践样本")
add_para(doc, "2025年以来，我们在新一代无人机航测平台（UAVPlatform）项目中试点了“AI人机协同”开发模式："
              "借助具备语义检索、多文件重构与Agent自主执行能力的AI编程助手进行迭代开发。截至目前，"
              "后端已产出231个Java文件、约1.63万行代码，前端产出60个Vue/TS文件、约1.33万行代码，"
              "并同步沉淀了44份需求、设计与排障文档，全部为人机协同完成。项目实践证明：AI模式在效率、"
              "质量与知识沉淀三个维度均显著优于传统模式，具备在部门全面推广的条件。")

# ==================== 二、建议内容 ====================
add_h1(doc, "二、建议内容")
add_para(doc, "总体建议：将“AI人机协同”升级为部门标准开发流程，建立“工具＋知识库＋规范＋度量”"
              "四位一体的落地机制。按四个阶段推进，每阶段给出具体做法与实施步骤。")

add_h2(doc, "（一）第一阶段：工具落地（第1～2周）")
add_bullet(doc, "1．统一为开发人员配备Agentic类AI编程助手（具备代码语义检索、跨文件编辑、Agent自主执行与验证能力），"
                "完成账号开通与环境配置；")
add_bullet(doc, "2．为每个在建项目建立AI工作区，导入代码仓库、接口文档与环境说明，使AI具备项目上下文；")
add_bullet(doc, "3．制定《AI辅助编码操作规程》并全员宣贯，明确三条红线：AI生成的代码必须经本地编译与测试验证后方可合入；"
                "AI对话产生的任务结论必须落库到本地工程文档；涉密项目代码严禁直接外发至公有云模型。")

add_h2(doc, "（二）第二阶段：知识库与AI记忆建设（第3～6周）")
add_bullet(doc, "1．将SDK文档、错误码字典、历史踩坑案例结构化整理后导入AI知识库。存量资产可直接复用："
                "本方向已积累44份技术文档与87个诊断脚本案例，均可作为首批入库语料；")
add_bullet(doc, "2．建立项目级“AI记忆”：固化编码规范、架构分层约定、UI样式规范、环境配置要点，"
                "使AI输出自动对齐团队标准，减少返工；")
add_bullet(doc, "3．将排障方法论固化为AI可执行的技能，明确“复现→定位→修复→验证”四阶段RCA流程，"
                "要求每次排障由AI按受控实验思想自动生成验证脚本，替代人工逐版试错。")

add_h2(doc, "（三）第三阶段：开发流程再造（第7～12周）")
add_para(doc, "在软件生命周期各环节嵌入AI协同点，形成新的标准作业流程：")
add_bullet(doc, "需求与设计环节：由AI根据需求描述生成架构图、时序图、接口草案与数据库表设计初稿，"
                "研发负责人评审修订后定稿，设计文档产出时间预计缩短一半以上；")
add_bullet(doc, "编码环节：样板代码（CRUD、DTO、配置、表单页）100%交由AI生成，开发人员聚焦业务逻辑"
                "与边界条件；代码定位以语义检索取代人工逐文件翻找；")
add_bullet(doc, "测试环节：AI依据代码变更自动生成单元测试与集成测试脚本，覆盖率纳入合入门禁；")
add_bullet(doc, "排障环节：AI按RCA流程自动生成受控实验脚本并解读日志，形成证据链后再动手修改，"
                "避免“盲改试错”；")
add_bullet(doc, "文档环节：接口文档、部署手册随代码变更由AI同步起草，人工审校，保证文档与代码不脱节。")

add_h2(doc, "（四）第四阶段：度量与推广（持续推进）")
add_bullet(doc, "1．建立效率度量看板，跟踪五项核心指标：需求交付周期、人均代码产出、缺陷密度、"
                "线上问题平均修复时长（MTTR）、代码评审一次通过率；")
add_bullet(doc, "2．每双周召开一次AI实践复盘会，迭代提示词模板与知识库内容，沉淀最佳实践；")
add_bullet(doc, "3．先在1～2个项目组试点一个月，形成可量化的对比案例后，再向全部门推广。")

# ==================== 三、可行性分析 ====================
add_h1(doc, "三、可行性分析")

add_h2(doc, "（一）技术可行性")
add_para(doc, "本建议已在UAVPlatform项目中完成实证：Spring Boot 3.3后端、Vue 3前端、DJI SDK深度集成、"
              "三维重建任务链路等复杂场景均可由AI高质量协同完成；同时，私有化开源大模型（如Ollama等）"
              "部署方案成熟，涉密场景可完全离线运行，技术路线不存在障碍。")

add_h2(doc, "（二）所需资源与投入成本")
add_table(doc,
          ["项目", "内容说明", "估算成本"],
          [["AI编程助手订阅", "按10人规模、约150元/人/月", "约1.8万元/年"],
           ["培训与导入", "4次专题培训＋导师结对辅导，内部消化", "约40人时（不产生直接费用）"],
           ["私有化模型（可选）", "复用现有GPU服务器＋开源模型，用于涉密项目", "0～3万元（一次性）"],
           ["知识库整理", "历史文档、诊断脚本结构化入库", "约20人天（分摊至各项目组）"],
           ["合计", "首年新增直接投入", "约2～5万元，无需新增编制"]],
          widths=[3.5, 8.0, 4.5])

add_h2(doc, "（三）风险点与应对思路")
add_table(doc,
          ["风险点", "潜在影响", "应对思路"],
          [["代码安全与保密", "代码外发公有云模型导致泄密", "涉密项目强制使用私有化部署模型；外发内容先行脱敏；操作规程划定红线并定期审计"],
           ["AI幻觉与低质量代码", "引入隐蔽缺陷，增加维护成本", "AI代码必须经本地编译、测试与人工评审后方可合入；关键模块以人工为主、AI为辅"],
           ["过度依赖", "团队基础编码能力退化", "保留编码基本功训练；核心算法与安全模块坚持人工主导设计"],
           ["版权与合规", "生成代码的商用合规风险", "统一采购正规商用订阅，开启合规检测选项"]],
          widths=[3.0, 4.5, 8.5])

add_h2(doc, "（四）组织可行性")
add_para(doc, "本建议不改变现有项目管理与交付流程，属于工具与作业方式的升级，实施阻力小；"
              "仅需指定1名兼职“AI工具负责人”统筹推进，各项目组按阶段计划执行即可，整体风险可控。")

# ==================== 四、预期效益 ====================
add_h1(doc, "四、预期效益")

add_h2(doc, "（一）经济效益")
add_para(doc, "参照本项目试点实测与传统模式的历史数据对比，预期各项指标改善如下：")
add_table(doc,
          ["对比维度", "传统编码模式（KOC2项目实测）", "AI协同模式（预期）", "改善幅度"],
          [["单功能开发周期", "5～10人天（含SDK调研试错）", "2～5人天", "缩短约50%"],
           ["线上问题排查时长", "1～3天/个（人均编写5～10版脚本）", "0.5天以内（AI生成验证脚本）", "缩短60%以上"],
           ["样板代码占用工时", "约占编码工作量60%", "趋近于0，由AI生成", "释放约一半编码工时"],
           ["新人上手周期", "1～2个月", "1～2周（AI记忆承载项目知识）", "缩短约75%"]],
          widths=[3.2, 5.6, 4.6, 2.6])
add_para(doc, "按10人研发团队、年人力成本约200万元测算，若整体研发效率提升40%，相当于每年释放约"
              "80万元人力价值；扣除首年工具与建设投入2～5万元，投入产出比超过15:1。同时交付周期缩短"
              "30%以上，意味着可以并行承接更多项目，直接增加合同营收。")

add_h2(doc, "（二）效率与质量提升")
add_bullet(doc, "单元测试覆盖率显著提升，缺陷左移到编码阶段暴露，线上故障率随之下降；")
add_bullet(doc, "接口文档、部署手册与代码同步更新，交付物完整度与规范性提高，减少验收阶段的文档返工；")
add_bullet(doc, "排障由“个人经验驱动”转为“证据链驱动”，问题闭环周期大幅缩短，运维压力减轻。")

add_h2(doc, "（三）满意度与组织能力建设")
add_bullet(doc, "开发人员从重复性“搬运代码”中解放出来，聚焦架构设计与业务创新，岗位价值感与工作满意度提升，"
                "有利于团队稳定；")
add_bullet(doc, "87个诊断脚本、44份技术文档等历史资产将转化为组织级知识库，核心经验不再依附于个人，"
                "有效对冲人员流动风险；")
add_bullet(doc, "形成可对外展示的“AI驱动开发”方法论与实证案例，提升部门技术品牌影响力，为承接更高复杂度"
                "项目奠定基础。")

# 落款
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(24)
r = p.add_run("建议人：×××")
set_font(r, "宋体", "Times New Roman", 12)
p = doc.add_paragraph()
r = p.add_run("2026年8月27日")
set_font(r, "宋体", "Times New Roman", 12)

out = r"d:\AiFlow\关于推行AI人机协同开发流程的建议书.docx"
doc.save(out)
print("saved:", out)
