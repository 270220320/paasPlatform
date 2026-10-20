# -*- coding: utf-8 -*-
"""生成《申请书》转正申请书 Word 文档（可打印格式）。

格式要求：
- 标题"申请书"：居中、加粗（二号 22pt）
- 正文：宋体小四（12pt）、首行缩进两字符、1.5 倍行距
- 落款（申请人/日期）：右对齐
- 保留入职日期、岗位、姓名等占位符
"""
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn

OUT_PATH = r"d:\申请书.docx"

FONT_BODY = "宋体"
SIZE_BODY = Pt(12)      # 小四
INDENT_2CH = Pt(24)     # 首行缩进两字符（小四 12pt * 2）


def set_cn_font(run, name=FONT_BODY, size=SIZE_BODY, bold=False):
    """同时设置西文与中文字体（python-docx 需单独设置 eastAsia）。"""
    run.font.name = name
    run.font.size = size
    run.font.bold = bold
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), name)


def add_para(doc, text, *, align=None, indent=True, bold=False,
             size=SIZE_BODY, font=FONT_BODY, space_after=Pt(0)):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    if indent:
        p.paragraph_format.first_line_indent = INDENT_2CH
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    p.paragraph_format.space_after = space_after
    run = p.add_run(text)
    set_cn_font(run, font, size, bold)
    return p


doc = Document()

# A4 纸张 + 常规页边距，便于直接打印
section = doc.sections[0]
section.page_width = Cm(21.0)
section.page_height = Cm(29.7)
section.top_margin = Cm(2.54)
section.bottom_margin = Cm(2.54)
section.left_margin = Cm(3.18)
section.right_margin = Cm(3.18)

# 标题：申请书（居中、加粗、二号）
add_para(doc, "申请书",
         align=WD_ALIGN_PARAGRAPH.CENTER, indent=False,
         bold=True, size=Pt(22), space_after=Pt(18))

# 称呼：顶格
add_para(doc, "尊敬的公司领导：", indent=False)

# 正文
add_para(doc, "您好！我于 XXXX 年 X 月 X 日入职公司，担任 XX 岗位，至今已满三个月。"
              "现试用期即将届满，特此提交转正申请，恳请公司予以批准，并自转正之日起按全额标准发放工资。"
              "三个月来的工作情况汇报如下：")

add_para(doc, "一、主要工作成果", bold=True)

add_para(doc, "1. 独立完成无人机管理平台（AiFlow）的研发落地。三个月来，我采用“AI 协同开发”模式，"
              "完成了该平台从需求分析、产品设计、原型验证到编码实现、测试部署的全流程研发，主要功能模块包括：")

add_para(doc, "设备管理：对接大疆司空（FlightHub 2）SDK，实现机场、无人机设备的接入、状态监控与远程调试；")
add_para(doc, "航线管理：集成航线编辑器，支持航线的在线创建、编辑、导入导出与飞行任务的自动下发执行；")
add_para(doc, "媒体管理与三维重建：实现飞行媒体文件自动归集、红外照片自动清理，并对接三维重建能力，有效保障重建成功率；")
add_para(doc, "平台采用前后端分离架构（Java Spring Boot + Vue 3），用户仅需浏览器即可使用，跨平台、免安装、便于统一部署维护。")

add_para(doc, "2. 推动 AI 研发流程在公司落地。在项目研发过程中，我沉淀了一套可复用的 AI 协同研发标准流程，"
              "覆盖需求确认、PRD 生成、原型设计、数据库设计、详细设计、编码与测试各环节，并通过项目实践验证了其效率与质量："
              "显著压缩了从需求到交付的周期；同步产出完整的文档体系，保证项目的可维护性与可交接性；"
              "已在公司内部进行实践汇报与分享，为团队后续项目采用 AI 研发方式提供了可参考的样板。")

add_para(doc, "二、个人情况说明", bold=True)

add_para(doc, "目前我个人背负住房按揭贷款，每月还款压力较大，经济负担较重。"
              "恳请公司领导体谅我的实际困难，自转正之日起按全额工资标准发放薪资，"
              "帮助我缓解还贷压力，使我能够更加安心、专注地投入工作。")

add_para(doc, "三、后续工作展望", bold=True)

add_para(doc, "如获批准转正，我将继续以 AI 协同研发方式投入工作，持续完善平台功能、保障系统稳定运行，"
              "并推动 AI 研发流程在更多项目中复制推广，为公司创造更大价值。")

add_para(doc, "恳请领导审阅批准！")

# 结尾：此致 / 敬礼
add_para(doc, "此致")
add_para(doc, "敬礼！", indent=False)

# 落款：右对齐
add_para(doc, "申请人：XXX", align=WD_ALIGN_PARAGRAPH.RIGHT, indent=False)
add_para(doc, "XXXX 年 XX 月 XX 日", align=WD_ALIGN_PARAGRAPH.RIGHT, indent=False)

doc.save(OUT_PATH)
print("saved:", OUT_PATH)
