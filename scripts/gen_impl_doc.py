# -*- coding: utf-8 -*-
"""生成《AiFlow 项目 AI 人机协同开发实施过程总结》Word 文档（融合 tkxm-sop 全流程，约1800字）。"""
import re
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn

OUT = r"d:\AiFlow\AiFlow项目AI协同开发实施过程总结.docx"

doc = Document()

def set_cn_font(run, east="宋体", ascii_font="Times New Roman", size=12, bold=False):
    run.font.name = ascii_font
    run.font.size = Pt(size)
    run.font.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), east)

def add_title(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    set_cn_font(r, east="黑体", size=16, bold=True)
    p.paragraph_format.space_after = Pt(12)

def add_h1(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_cn_font(r, east="黑体", size=14, bold=True)
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(6)

def add_para(text, indent=True):
    p = doc.add_paragraph()
    if indent:
        p.paragraph_format.first_line_indent = Cm(0.74)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    p.paragraph_format.line_spacing = 1.3
    r = p.add_run(text)
    set_cn_font(r)
    return p

add_title("AiFlow（UAVPlatform）无人机作业平台\nAI 人机协同开发实施过程总结")

add_h1("一、项目概述")
add_para("AiFlow（产品定名 UAVPlatform）是一套面向固定场地、单一作业主体的无人机云端作业平台，与私有化部署的 DJI 司空 2（FlightHub 2 PaaS，192.168.76.141）深度集成：航线编辑、设备管理、驾驶舱直播与飞行回放复用司空前端组件 paas.js；任务下发、设备/航线镜像同步、限飞区推送与三维重建则通过 FlightHub Java SDK 与 REST 网关完成。平台功能覆盖数据看板、地图标注、空域安全、航线库、任务计划、媒体库与三维重建、人员权限与设备管理等模块。项目由 1 名研发协同 AI 实施约 4 周，完成后端 231 个 Java 文件（约 1.63 万行）、前端 60 个文件（约 1.33 万行）及 44 份设计与排障文档。")

add_h1("二、实施提纲：以 tkxm-sop 驱动全流程流转")
add_para("项目以自研的 tkxm-sop 研发 SOP 技能套件为总编排，实现从需求到测试的全流程流转：需求 → 产品（PRD/原型/用户计划）→ 设计（概要/数据库/详细/开发计划）→ 开发（编码/评审）→ 测试，共 5 个阶段、11 个流程节点，另配初始化与变更处理两个横切流程。整套 SOP 坚持三条原则：产物驱动而非记忆驱动——扫描 docs/ 目录已有产物即可推断当前所处阶段；门控推进——上游产物未经确认不进入下游；可跳步但须声明缺失与风险。人负责架构决策与验收拍板，AI 负责各环节文档起草与代码生成，产出统一落盘 docs/ 形成可回溯基线。")

add_h1("三、所用工具与模型")
add_para("（1）AI IDE：Qoder 作为研发主界面，提供仓库级语义检索、代码智能（跳转/引用分析）、Agent 自主执行与规则记忆机制；复杂任务先经计划模式产出实施方案、由代码评审子代理复核改动质量，使产出自动对齐团队规范。", indent=False)
add_para("（2）模型层：采用 GLM-5.2（智谱）国产合规组合，中文理解与代码生成能力强，国内网络直连稳定；涉密场景可退回 Ollama 承载的私有化开源模型，保证数据不出境。", indent=False)
add_para("（3）MCP 工具：部署 FlightHub MCP Server（fh_* 工具集），Agent 可直接查询设备列表、上传并执行航线、下发飞控指令、获取遥测数据，实现“对话即运维”。", indent=False)
add_para("（4）工程工具链：后端 JDK 21 + Maven + Spring Boot 3.3 + MyBatis Plus + MySQL 8；前端 Node 18 + Vite + Vue3 + Element Plus；WSL 承载后端运行环境，配套 87 个诊断排障脚本。", indent=False)

add_h1("四、两层 Agent Skill 体系")
add_para("全流程层——tkxm-sop 套件（13 个 skill）：tkxm-init 一键安装套件并搭建 docs/ 产物目录骨架；01～08 号流程依次产出需求说明书、PRD、原型图、用户计划、概要设计（DDD 方法论）、数据库设计、详细设计（TDD/SDD）与开发计划；09 编码并同步单测、10 代码审查、11 生成测试计划/用例/冒烟方案；tkxm-change 处理定稿后的变更——沿依赖图计算“爆炸半径”，按拓扑序增量回灌下游产物并写变更记录留痕。")
add_para("领域层——cloudapi2 系列：sdk-core（FlightHub Java SDK 集成规范：初始化、飞控权限、HMS、任务下发）、frontend-component（paas.js 七大前端组件接入指南）、mcp-guide（fh_* 工具选择与编排）、troubleshooting（四阶段 RCA 与错误码字典）等，为司空集成提供专业领域知识。")

add_h1("五、实施步骤：从需求到编码")
add_para("1. 套件初始化：tkxm-init 将 13 个 skill 安装到工作区并校验完整性，创建 docs/01-requirements～05-test 产物目录骨架，输出五阶段流程总览与起步入口。", indent=False)
add_para("2. 需求阶段：以司空 2 界面截图与业务对话为输入，tkxm-requirement 通过头脑风暴梳理并聚焦核心功能点，产出需求说明书（如三维重建专项需求）落盘 docs/01-requirements/。", indent=False)
add_para("3. 产品阶段：tkxm-prd 生成 PRD 并经确认门控迭代至 V1.0 定稿，确立“单作业空间、不做多项目”的裁剪原则与 P0/P1/P2 优先级；tkxm-prototype 按约定的 Element Plus 深色主题产出高保真可点选原型；tkxm-userplan 与需求方沟通确定功能点迭代顺序。", indent=False)
add_para("4. 设计阶段：tkxm-preliminary 以 DDD 方法产出概要设计（模块化单体、11 个限界上下文分包）与全套 ER 图；tkxm-database 产出数据库设计与建表语句（ddl-v1.0.sql 及司空增量 ddl-fh2-v1.1.sql）；tkxm-detail 分模块产出详细设计；tkxm-devplan 制定开发计划，将功能点统一编号为 U-001～U-091 并按依赖关系划分波次。关键架构决策由人拍板。", indent=False)
add_para("5. 开发阶段：先经骨架部署打通司空鉴权（user_id/JWT）、EMQX、Webhook 等环境配置；随后 tkxm-coding 按开发计划波次推进——后端遵循阿里巴巴 Java 开发手册，按 iam/device/map/safety/wayline/mission/media/fh2 分层分包，以反射方式适配 FlightHub SDK 完成设备、航线同步与任务下发，SQL 遵循腾讯开发规范；前端接入 WaylineEditor 航线编辑器、设备管理、驾驶舱等 paas.js 组件，由 cloudapi2 系列 skill 提供接入规范。每个功能点完成即编写单元测试，核心逻辑覆盖率不低于 80%。", indent=False)
add_para("6. 评审与测试：tkxm-review 对照概要、详细与数据库设计审查代码，产出检查报告并回流修改；tkxm-test 基于原型与设计文档生成测试计划、测试方案、用例集与冒烟方案。联调中的环境与 SDK 兼容性问题，借助 MCP 工具、诊断脚本与四阶段 RCA 闭环，典型缺陷均形成排障报告归档。", indent=False)
add_para("7. 变更闭环：开发中需求或设计一旦变动，tkxm-change 从变更入口沿依赖图计算爆炸半径，按拓扑序增量回灌受影响的 PRD、设计、代码与用例，并写 C 编号变更记录留痕，保证全链路一致。", indent=False)

add_h1("六、成效与小结")
add_para("4 周内 1 人完成全平台交付，样板代码工时趋近于零；文档随代码同步演进，从负担变为资产；踩坑经验沉淀为知识库与 Skill。自研 tkxm-sop 把研发流程固化为 Agent 可执行的技能，实现“流程即代码”，使全流程流转不依赖个人经验。实践心得：输入决定输出，需求与边界描述越清晰、产出越准；架构与安全决策必须由人拍板；编译、测试、人工评审三道关是控制 AI 幻觉、保障质量的关键。")

doc.save(OUT)

# 统计正文字数（Word 口径：中文字符 + 中文标点 + 英文单词）
text = "\n".join(p.text for p in doc.paragraphs)
cn = len(re.findall(r"[\u4e00-\u9fff]", text))
cn_punc = len(re.findall(r"[\u3000-\u303f\uff00-\uffef]", text))
en_words = len(re.findall(r"[A-Za-z0-9][A-Za-z0-9./_-]*", text))
total = len(re.sub(r"\s", "", text))
print(f"saved: {OUT}")
print(f"paragraphs: {len(doc.paragraphs)}  cn: {cn}  cn_punc: {cn_punc}  en_words: {en_words}")
print(f"Word-count(est): {cn + cn_punc + en_words}  total-chars(no-space): {total}")
