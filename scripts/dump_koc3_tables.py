# -*- coding: utf-8 -*-
"""打印 KOC3 报告中两个目标表格的详细内容"""
import sys
from docx import Document

sys.stdout.reconfigure(encoding='utf-8')

path = r"D:\KOC3\大疆上云 API 2 KOC 3 测试报告模板_20260826(1).docx"
doc = Document(path)

for ti, table in enumerate(doc.tables):
    hdr = ' | '.join(c.text.strip() for c in table.rows[0].cells)
    print(f"=== Table {ti}: {len(table.rows)}rows x {len(table.columns)}cols | header: {hdr}")
    if '是否解决' in hdr or '需求标题' in hdr:
        for ri, row in enumerate(table.rows):
            cells = [c.text.strip().replace('\n', '\\n') for c in row.cells]
            print(f"  row{ri}: {cells}")
        # 检查数据行单元格字体格式
        from docx.shared import Pt
        c = table.rows[1].cells[1]
        for p in c.paragraphs:
            for r in p.runs:
                print(f"  sample run font: name={r.font.name}, size={r.font.size}, bold={r.font.bold}")
