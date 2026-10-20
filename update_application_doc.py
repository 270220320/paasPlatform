from copy import deepcopy
from docx import Document
from pathlib import Path

path = Path(r'D:\AiFlow\全薪申请书_孟子文_待修改.docx')
doc = Document(path)

# Correct the summary so it reflects the actual number of entries in the project feature register.
p = doc.paragraphs[2]
old = '其中，上云2平台已完成42项功能点，交管0.5累计UI优化页面32个。'
new = '其中，上云2平台功能清单梳理了42项功能点，交管0.5版本累计优化32个页面。'
full = p.text
if old not in full:
    raise RuntimeError('Expected summary text was not found; leaving the document unchanged.')
updated = full.replace(old, new)

def set_paragraph_text(paragraph, text):
    old_runs = list(paragraph.runs)
    rpr = deepcopy(old_runs[0]._r.rPr) if old_runs and old_runs[0]._r.rPr is not None else None
    for run in old_runs:
        paragraph._p.remove(run._r)
    run = paragraph.add_run(text)
    if rpr is not None:
        run._r.insert(0, rpr)

set_paragraph_text(p, updated)

set_paragraph_text(doc.paragraphs[5], '1. 上云2无人机管理平台研发。在不到两个月的时间里，我结合大疆上云API2，借助AI Agent独立推进平台研发，完成需求梳理、产品原型设计、数据库设计、前后端开发和自动化测试等工作。项目功能清单共梳理42项，涉及账号与权限、设备、航线、任务计划、地图标注、空域安全、媒体管理、驾驶舱和飞行回放等功能。')
set_paragraph_text(doc.paragraphs[6], '在具体功能实现中，我对接大疆上云API2 SDK，完成机场与无人机设备接入、状态查看和远程调试；集成航线编辑器，支持航线创建、编辑、导入导出，并衔接飞行任务下发；同时建设首页数据总览、飞行任务管理、安全空域、地图标注、媒体管理及人员权限等能力。')
set_paragraph_text(doc.paragraphs[7], '2. 梳理并实践AI协同研发流程。在平台研发过程中，我将AI工具应用于需求确认、PRD编写、原型与数据库设计、详细设计、编码和测试，并记录实践中的经验与问题，形成AI研发说明和经验文档，为团队后续项目开展AI协同研发提供参考。')
set_paragraph_text(doc.paragraphs[8], '3. 交管产品0.5版本UI优化。按照UI设计稿，对累计32个页面进行高保真还原，并整理复用组件、统一色值变量，改善页面呈现并提升后续维护的一致性。此外，我参与了公司金点子活动。')

doc.save(path)
